#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Tests autonomes de l'export « dossier d'analyse complet » (`tools/export/dossier.py`) — aucun modèle, aucun réseau."""
from __future__ import annotations

import json
import sqlite3
import sys
import tempfile
import zipfile
from pathlib import Path

RACINE = Path(__file__).resolve().parents[3]
if str(RACINE) not in sys.path:
    sys.path.insert(0, str(RACINE))

from tools.export import dossier as dos  # noqa: E402
from tools.export import export as export_tool  # noqa: E402

ECHECS: list[str] = []
NL = chr(10)


def verifier(nom: str, condition: bool) -> None:
    print(("PASS " if condition else "FAIL ") + nom)
    if not condition:
        ECHECS.append(nom)


D = Path(tempfile.mkdtemp(prefix="e21-dossier-"))
(D / "intrants").mkdir()
(D / "intrants" / "secret.md").write_text("DOCUMENT D'ORIGINE CONFIDENTIEL", encoding="utf-8")
(D / "connaissance").mkdir()
base = sqlite3.connect(D / "connaissance" / "index.sqlite")
base.execute("CREATE TABLE elements (id INTEGER PRIMARY KEY, genre TEXT, doc TEXT, libelle TEXT, extrait TEXT, verifie INTEGER)")
base.execute("INSERT INTO elements VALUES (12, 'fait', 'Audit_interne.docx.md', 'MFA absente', 'Deux comptes admin partagés sans MFA', 1)")
base.execute("INSERT INTO elements VALUES (34, 'fait', 'Inventaire.md.md', 'Sauvegarde', 'Sauvegarde quotidienne sur NAS non testée', 1)")
base.commit()
base.close()


def ecrire(nom: str, lignes: list[str]) -> None:
    (D / nom).write_text(NL.join(lignes) + NL, encoding="utf-8")


ecrire("01-actifs.md", ["# Actifs", "", "| ID | Actif | Type | Description | Sensibilité | Preuve source |", "|---|---|---|---|---|---|",
                        "| A01 | **NordTMS** | Logiciel | ERP de transport <script>alert(1)</script> [E12] | Critique | [E12] |",
                        "| A02 | NAS local | Matériel | Sauvegardes [E34] | Interne | [E34] |"])
ecrire("03-menaces.md", ["# Menaces", "", "| ID | Actif concerné | Catégorie | Description | CVE | Source |", "|---|---|---|---|---|---|",
                         "| M001 | A01 — NordTMS | STRIDE-S | Usurpation d'un compte admin partagé [E12] | — | STRIDE-S |",
                         "| M002 | A02 — NAS local | STRIDE-D | Perte des sauvegardes non testées [E34] [E999] | — | STRIDE-D |"])
ecrire("04-evaluation.md", ["# Évaluation", "", "| ID | Menace | Probabilité | Impact | Niveau | Justification |", "|---|---|---|---|---|---|",
                            "| M001 | A01 STRIDE-S | Élevé | Élevé | **Critique** | MFA absente [E12] |",
                            "| M002 | A02 STRIDE-D | Moyen | Élevé | Élevé | Test de restauration jamais fait [E34] |"])
ecrire("05-traitement.md", ["# Traitement", "", "| ID | Actif et menace | Traitement | Contre-mesures | Sources | Risque résiduel estimé |", "|---|---|---|---|---|---|",
                            "| M001 | A01 STRIDE-S | Réduire | MFA obligatoire sur les comptes admin [E12] | ISO27002-5.15 | Moyen |",
                            "| M002 | A02 STRIDE-D | Réduire | Tests de restauration trimestriels | ISO27002-8.13 | Faible |"])
ecrire("registre-risques.md", ["# Registre", "", "| ID | Actif | Menace | Niveau | Traitement | valide_par | Statut |", "|---|---|---|---|---|---|---|",
                               "| M001 | A01 — NordTMS | STRIDE-S compte partagé | Critique | Réduire | Marie · accepté · 2026-10-12 | Décidé |",
                               "| M002 | A02 — NAS local | STRIDE-D perte des sauvegardes | Élevé | Réduire |  | À valider |"])
ecrire("SYNTHESE.md", ["# Synthèse", "", "Deux risques : M001 (critique) et M002.", "", "- Voir A01 et [E12]."])
(D / "registre_risques.json").write_text(json.dumps({"risques": [{"id": "M001"}]}), encoding="utf-8")
(D / "validations.json").write_text(json.dumps({"decisions": {"x": {
    "decision": "accepte", "libelle_decision": "Accepté tel quel", "analyste": "Marie", "date": "2026-10-12 10:00",
    "type": "risque", "titre": "M001 — NordTMS", "commentaire": "OK pour moi"}}}), encoding="utf-8")

# --- HTML : liens, fiches, preuves, décisions -------------------------------------------------------------------------------------------
h = dos.html_complet(D)
verifier("un seul fichier HTML autonome, sommaire présent", h.startswith("<!DOCTYPE html>") and 'class="sommaire"' in h and "<style>" in h)
verifier("chaque livrable présent est une section du dossier", all(f'id="doc-{a}"' in h for a in ("synthese", "registre", "actifs", "menaces", "evaluation", "traitement")))
verifier("le contenu des documents est échappé (aucune balise exécutable)", "<script>alert(1)</script>" not in h and "&lt;script&gt;" in h)
verifier("ancres de ligne : chaque actif et chaque risque a un identifiant de lien", 'id="actifs-A01"' in h and 'id="menaces-M001"' in h
         and 'id="evaluation-M002"' in h and 'id="traitement-M001"' in h)
verifier("l'identifiant d'un risque dans un tableau renvoie à sa fiche", 'href="#fiche-M001"' in h and 'id="fiche-M001"' in h)
verifier("l'actif cité dans une menace renvoie à sa ligne", 'href="#actifs-A01"' in h)
verifier("la fiche relie actif, menace, évaluation, traitement, validation et preuves",
         "Usurpation d'un compte admin partagé" in h and "Probabilité : Élevé" in h and "MFA obligatoire" in h and "Marie · accepté" in h)
verifier("la fiche renvoie à chaque ligne d'origine", 'href="#menaces-M001"' in h and 'href="#evaluation-M001"' in h and 'href="#registre-M001"' in h)
verifier("les marques de gras ne fuient pas dans les fiches", "**Critique**" not in h.split('id="fiche-M001"')[1].split("</div>")[0])
verifier("chaque preuve [E12] renvoie à l'annexe avec son document et son extrait exact", 'href="#preuve-12"' in h
         and 'id="preuve-12"' in h and "Audit_interne.docx.md" in h and "Deux comptes admin partagés sans MFA" in h)
verifier("l'annexe dit quels risques et actifs s'appuient sur la preuve", "Cité par" in h.split('id="preuve-12"')[1][:600] and 'href="#fiche-M001"' in h.split('id="preuve-12"')[1][:900])
verifier("une preuve introuvable est signalée, jamais inventée", 'class="preuve inconnue"' in h and "E999" in h and "introuvable" in h)
verifier("historique des décisions : qui, quoi, quand, motif", "Historique des décisions" in h and "OK pour moi" in h and "Accepté tel quel" in h)
verifier("les documents d'origine ne sont pas inclus", "CONFIDENTIEL" not in h)

# --- graphe de traçabilité ----------------------------------------------------------------------------------------------------------------
g = dos.graphe(D)
types = {n["type"] for n in g["noeuds"]}
liens = {(l["source"], l["cible"], l["type"]) for l in g["liens"]}
verifier("graphe : documents, preuves, actifs et risques", {"document", "preuve", "actif", "risque"} <= types)
verifier("graphe : risque → actif, risque → preuve, preuve → document", ("M001", "A01", "concerne") in liens
         and ("M001", "E12", "s_appuie_sur") in liens and ("E12", "doc:Audit_interne.docx.md", "extrait_de") in liens)

# --- tableur et CSV ---------------------------------------------------------------------------------------------------------------------------
csv = dos.csv_registre(D)
verifier("CSV : BOM, séparateur « ; », une ligne par risque, preuves et documents", csv.startswith(chr(0xFEFF)) and ";" in csv.splitlines()[0]
         and "valide_par" in csv.splitlines()[0]
         and len(csv.strip().splitlines()) == 3 and "Audit_interne.docx.md" in csv)
x = dos.xlsx_registre(D, D / "t.xlsx")
if x is not None:
    from openpyxl import load_workbook
    wb = load_workbook(x)
    verifier("XLSX : feuilles Registre, Actifs, Preuves, Liens", wb.sheetnames == ["Registre", "Actifs", "Preuves", "Liens"]
             and wb["Registre"].max_row == 3)
    x.unlink()

# --- archive ---------------------------------------------------------------------------------------------------------------------------------
sortie = Path(tempfile.mkdtemp(prefix="e21-dossier-out-"))
z = dos.exporter_dossier(str(D), "zip", str(sortie))
noms = zipfile.ZipFile(z).namelist()
verifier("zip : HTML, livrables, registre tableur et CSV, graphe, décisions, mode d'emploi",
         {"dossier-complet.html", "LISEZ-MOI.txt", "livrables/03-menaces.md", "livrables/SYNTHESE.md", "donnees/registre-risques.csv",
          "donnees/graphe-tracabilite.json", "donnees/decisions.json", "donnees/registre_risques.json"} <= set(noms))
verifier("zip : aucun document d'origine ni index interne", not any("intrants" in n or "secret" in n or n.endswith(".sqlite") for n in noms))
verifier("zip : aucun fichier temporaire laissé", not [p for p in sortie.iterdir() if p.name.startswith(".tmp")])
verifier("zip : le mode d'emploi décrit la chaîne de liaison", "document → preuve" in zipfile.ZipFile(z).read("LISEZ-MOI.txt").decode("utf-8"))

# --- par le point d'entrée habituel de l'export -----------------------------------------------------------------------------------------------
sortie2 = Path(tempfile.mkdtemp(prefix="e21-dossier-out2-"))
chemin = export_tool.exporter(str(D), "dossier-html", out=str(sortie2))
verifier("export.exporter(« dossier-html ») fonctionne et écrit dans le dossier demandé", chemin.parent == sortie2.resolve() and chemin.is_file())
verifier("l'export exécutif d'avant reste inchangé (md)", export_tool.exporter(str(D), "md", out=str(sortie2)).name.endswith(".md"))

# --- dossier presque vide : jamais de plantage ----------------------------------------------------------------------------------------------------
V = Path(tempfile.mkdtemp(prefix="e21-dossier-vide-"))
(V / "SYNTHESE.md").write_text("# Synthèse\n\nRien d'autre.\n", encoding="utf-8")
hv = dos.html_complet(V)
verifier("dossier presque vide : HTML produit, sans fiche ni preuve", "Rien d'autre" in hv and "annexe-fiches" not in hv)
try:
    dos.exporter_dossier(str(V / "inexistant"), "html")
    verifier("dossier inexistant : refusé proprement", False)
except ValueError:
    verifier("dossier inexistant : refusé proprement", True)

print("EXPORT-DOSSIER:", "OK" if not ECHECS else f"{len(ECHECS)} ECHEC(S)")
sys.exit(1 if ECHECS else 0)

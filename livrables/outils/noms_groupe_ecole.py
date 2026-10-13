"""Ajoute les noms du groupe sur la SLIDE DE TITRE du PPT école (seule modification ; le reste du fichier est inchangé)."""
import copy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

source, sortie = Path(sys.argv[1]), Path(sys.argv[2])
prs = Presentation(str(source))
cover = prs.slides[0]
nom = date = None
for sh in cover.shapes:
    if sh.has_text_frame:
        t = sh.text_frame.text.strip()
        if t == "Melvin RAIMBAULT":
            nom = sh
        elif t.startswith("Soutenance"):
            date = sh
assert nom is not None and date is not None, "zones de la couverture introuvables"

# deux lignes de noms, même police que « Melvin RAIMBAULT »
p0 = nom.text_frame.paragraphs[0]
p0.runs[0].text = "Melvin RAIMBAULT  ·  Lucas HUARD"
for r in p0.runs[1:]:
    r.text = ""
p1 = copy.deepcopy(p0._p)
p0._p.addnext(p1)
p_nouveau = nom.text_frame.paragraphs[1]
p_nouveau.runs[0].text = "Jérémy LESSA  ·  David AMARO MONTEIRO"
nom.height = Inches(0.6)
date.top = Inches(4.98)
prs.save(str(sortie))
print("ok ->", sortie)

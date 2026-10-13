"""Aperçu approximatif d'un PPTX (Pillow + polices Windows) : planches de 6 slides + détection des débordements de texte."""
import io
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

PPTX, SORTIE = Path(sys.argv[1]), Path(sys.argv[2])
SORTIE.mkdir(exist_ok=True)
DPI = 144
FONTS = {"Segoe UI": "segoeui.ttf", "Segoe UI Semibold": "seguisb.ttf", "Consolas": "consola.ttf"}
cache = {}


def police(nom, pt, gras=False):
    fichier = FONTS.get(nom or "Segoe UI", "segoeui.ttf")
    if gras and nom == "Segoe UI":
        fichier = "segoeuib.ttf"
    cle = (fichier, round(pt * DPI / 72))
    if cle not in cache:
        cache[cle] = ImageFont.truetype("C:/Windows/Fonts/" + fichier, cle[1])
    return cache[cle]


def px(emu):
    return int(Emu(emu).inches * DPI)


def hexa(c, defaut="000000"):
    return "#" + (str(c) if c else defaut)


prs = Presentation(str(PPTX))
L, H = int(Emu(prs.slide_width).inches * DPI), int(Emu(prs.slide_height).inches * DPI)
problemes = []
vignettes = []
for n, slide in enumerate(prs.slides, 1):
    try:
        fond = hexa(slide.background.fill.fore_color.rgb, "FFFFFF")
    except Exception:
        fond = "#FFFFFF"
    im = Image.new("RGB", (L, H), fond)
    d = ImageDraw.Draw(im)
    for sh in slide.shapes:
        x, y, w, h = px(sh.left), px(sh.top), px(sh.width), px(sh.height)
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                p = Image.open(io.BytesIO(sh.image.blob)).convert("RGBA")
                p = p.resize((max(w, 1), max(h, 1)))
                im.paste(p, (x, y), p)
            except Exception:
                d.rectangle([x, y, x + w, y + h], outline="#999999")
            continue
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            try:
                f = hexa(sh.fill.fore_color.rgb) if sh.fill.type == 1 else None
            except Exception:
                f = None
            try:
                o = hexa(sh.line.color.rgb) if sh.line.fill.type == 1 else None
            except Exception:
                o = None
            try:
                prst = sh._element.spPr.prstGeom.get("prst")
            except Exception:
                prst = ""
            arrondi = prst == "roundRect"
            if prst == "line":
                d.line([x, y, x + w, y + h], fill=o or "#2DD4BF", width=2)
                continue
            if arrondi and w > 4 and h > 4:
                r = int(min(w, h) * (sh.adjustments[0] if len(sh.adjustments) else 0.16))
                d.rounded_rectangle([x, y, x + w, y + h], radius=r, fill=f, outline=o)
            elif f or o:
                d.rectangle([x, y, x + w, y + h], fill=f, outline=o)
        if sh.has_text_frame and sh.text_frame.text.strip():
            tf = sh.text_frame
            marge_x = px(tf.margin_left) if tf.margin_left is not None else 0
            marge_y = px(tf.margin_top) if tf.margin_top is not None else 0
            largeur = w - 2 * marge_x
            lignes = []  # (texte, police, couleur)
            for para in tf.paragraphs:
                runs = para.runs
                if not runs:
                    lignes.append(("", police("Segoe UI", 10), "#000000", None))
                    continue
                r0 = runs[0]
                pt = r0.font.size.pt if r0.font.size else 12
                pol = police(r0.font.name, pt, bool(r0.font.bold))
                try:
                    col = hexa(r0.font.color.rgb)
                except Exception:
                    col = "#000000"
                mots = "".join(r.text for r in runs).split(" ")
                cur = ""
                for m in mots:
                    essai = (cur + " " + m).strip()
                    if d.textlength(essai, font=pol) <= largeur or not cur:
                        cur = essai
                    else:
                        lignes.append((cur, pol, col, para.alignment))
                        cur = m
                lignes.append((cur, pol, col, para.alignment))
            hl = [int(p[1].size * 1.25) for p in lignes]
            total = sum(hl)
            ancre = tf.vertical_anchor
            yy = y + marge_y
            if ancre is not None and "MIDDLE" in str(ancre):
                yy = y + (h - total) // 2
            if total > h + int(0.08 * DPI) and sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
                problemes.append((n, tf.text.strip()[:50], round(total / DPI, 2), round(h / DPI, 2)))
            for (t, pol, col, al), lh in zip(lignes, hl):
                xx = x + marge_x
                if al is not None and "RIGHT" in str(al):
                    xx = x + w - marge_x - d.textlength(t, font=pol)
                elif al is not None and "CENTER" in str(al):
                    xx = x + (w - d.textlength(t, font=pol)) / 2
                d.text((xx, yy), t, font=pol, fill=col)
                yy += lh
    im.save(SORTIE / f"s{n:02d}.png")
    vignettes.append(im.resize((960, int(960 * H / L))))
# planches 2 × 3
for k in range(0, len(vignettes), 6):
    groupe = vignettes[k:k + 6]
    vh = groupe[0].size[1]
    planche = Image.new("RGB", (1920 + 10, (vh + 10) * 3), "#888888")
    for i, v in enumerate(groupe):
        planche.paste(v, ((i % 2) * 970, (i // 2) * (vh + 10)))
    planche.save(SORTIE / f"planche{k // 6 + 1}.png")
print("slides:", len(prs.slides))
for p in problemes:
    print("DÉBORDEMENT slide", p[0], "| besoin", p[2], "po > boîte", p[3], "po |", p[1])
